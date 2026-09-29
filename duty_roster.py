"""Read the supplied roster; keep edits in local JSON, never in the source XLSX."""
from __future__ import annotations

import calendar
import hashlib
import json
import os
import re
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

VIETNAM = ZoneInfo("Asia/Ho_Chi_Minh")


def normalize(value):
    return " ".join(unicodedata.normalize("NFC", str(value or "")).casefold().split())


def staff_key(value):
    """Return an opaque stable catalog ID without publishing employee codes."""
    return hashlib.sha256(("TOMS-SLA-MB:" + str(value)).encode("utf-8")).hexdigest()


def catalog_staff(rosters):
    people = {}
    for roster in rosters.values():
        for ident, person in roster.get("people", {}).items():
            staff_id = staff_key(ident)
            name = str(person.get("name", "")).strip()
            if name:
                people[staff_id] = {"staff_id": staff_id, "ten_nhan_su": name}
    return list(people.values())


def read_contacts(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    contacts = {}
    try:
        for sheet in book:
            rows = iter(sheet.iter_rows(values_only=True))
            for row in rows:
                labels = [normalize(v) for v in row]
                if "họ và tên" not in labels or "id telegram" not in labels:
                    continue
                name_col, id_col = labels.index("họ và tên"), labels.index("id telegram")
                for item in rows:
                    name, ident = item[name_col], item[id_col]
                    if not name or ident is None:
                        continue
                    if isinstance(ident, float) and ident.is_integer():
                        ident = int(ident)
                    ident = str(ident).strip()
                    if not re.fullmatch(r"[1-9]\d*", ident):
                        raise ValueError(f"ID Telegram không hợp lệ của {name}")
                    key = normalize(name)
                    if key in contacts and contacts[key] != ident:
                        raise ValueError(f"Tên trùng có ID Telegram khác nhau: {name}")
                    contacts[key] = ident
                break
        if not contacts:
            raise ValueError("Không tìm thấy cột Họ và tên / ID Telegram")
        return contacts
    finally:
        book.close()


def read_roster(path, year, month):
    days = calendar.monthrange(year, month)[1]
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "Bang ca" not in book.sheetnames:
            raise ValueError("File lịch phải có sheet Bang ca")
        rows = list(book["Bang ca"].iter_rows(values_only=True))
        title = " ".join(str(v) for row in rows[:7] for v in row if v is not None)
        period = re.search(r"Tháng\s+(\d{1,2})/(\d{4})", title, re.I)
        if not period or (int(period[2]), int(period[1])) != (year, month):
            raise ValueError("Tháng/năm chọn không khớp tiêu đề file lịch")
        header = next((i for i, row in enumerate(rows) if "họ và tên cbcnv" in [normalize(v) for v in row]), None)
        if header is None:
            raise ValueError("Thiếu cột Họ và tên CBCNV")
        labels = [normalize(v) for v in rows[header]]
        name_col, staff_col = labels.index("họ và tên cbcnv"), labels.index("mã nhân viên")
        day_row = next((i for i in range(header + 1, min(header + 5, len(rows))) if rows[i][4:7] == (1, 2, 3)), None)
        if day_row is None:
            raise ValueError("Không nhận diện được hàng ngày 1..31")
        columns = {int(v): i for i, v in enumerate(rows[day_row][4:], 4) if isinstance(v, int) and 1 <= v <= days}
        if set(columns) != set(range(1, days + 1)):
            raise ValueError("Lịch thiếu hoặc trùng ngày trong tháng")
        people = {}
        for row in rows[day_row + 1:]:
            ident = str(row[staff_col] or "").strip()
            if not ident.startswith("VNPT") or not row[name_col]:
                continue
            if ident in people:
                raise ValueError(f"Mã nhân viên trùng: {ident}")
            shifts = {str(day): str(row[col] or "").strip() for day, col in columns.items()}
            people[ident] = {"name": str(row[name_col]).strip(), "shifts": shifts}
        if not people:
            raise ValueError("Không có nhân sự trong lịch")
        return {"year": year, "month": month, "people": people}
    finally:
        book.close()


def load(path):
    try:
        result = json.loads(Path(path).read_text(encoding="utf-8"))
        return result if isinstance(result, dict) else {}
    except (OSError, ValueError):
        return {}


def save(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.chmod(0o600)
    os.replace(temporary, path)


def on_duty(rosters, now=None):
    now = now or datetime.now(VIETNAM)
    if now.tzinfo is not None:
        now = now.astimezone(VIETNAM)
    date = now.date() - timedelta(days=1) if now.hour < 8 else now.date()
    shift = "HC" if 8 <= now.hour < 17 else "Đêm"
    roster = rosters.get(date.strftime("%Y-%m"), {})
    return [dict(person, staff_id=staff_key(ident))
            for ident, person in roster.get("people", {}).items()
            if normalize(person["shifts"].get(str(date.day))) == normalize(shift)]


def recipients(rosters, contacts, now=None, staff_ids=None):
    people = on_duty(rosters, now)
    if staff_ids is not None:
        allowed = set(map(str, staff_ids))
        people = [person for person in people if person["staff_id"] in allowed]
        if not people:
            return []
    missing = [p["name"] for p in people if normalize(p["name"]) not in contacts]
    if not people:
        raise ValueError("Chưa có lịch/nhân sự cho ca trực hiện tại")
    if missing:
        raise ValueError("Thiếu ID Telegram: " + ", ".join(missing))
    return list(dict.fromkeys(contacts[normalize(p["name"])] for p in people))
